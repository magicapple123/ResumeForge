# SPDX-License-Identifier: Apache-2.0
"""Word 的局部排版落地（Q11）：缩进 / 对齐 / 加粗 / 斜体 / 字距 / 日期位置。

从 test_docx_exporter.py 拆出（2026-10-31，纯搬迁）：原文件超过 500 行预算；
本组用例只守 W07/W3 的五项局部排版在 .docx 里的投影（对象层 / 对象角色层）。
"""

from io import BytesIO

import pytest
from app.schemas.resume import ResumeContent
from app.services.docx_exporter import build_resume_docx
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm

# ===== Q11：局部排版属性在 Word 里的落地（缩进 / 对齐 / 加粗 / 斜体 / 字距）=====
#
# 2026-10-31 用户口径：同一份局部排版在 HTML / Word / PDF 三个出口必须同效。此前 Word
# 固定用模板 li_indent、run 不吃 style.bold/italic、无字距、正文段落不吃对齐——面板里
# 调了，导出后纸面不变。

_STYLE_ENTRY_ID = "exp_12345678"


def _styled_resume() -> ResumeContent:
    """最小真实简历：一条经历，id 与 format_config.objects / object_roles 的键对齐。"""
    return ResumeContent.model_validate(
        {
            "photo": "",
            "name": "测试",
            "gender": "",
            "birth_year": "",
            "phone": "",
            "email": "",
            "city": "",
            "personal_website": "",
            "github": "",
            "job_intent": "工程师",
            "summary": "",
            "education": [],
            "experience": [
                {
                    "id": _STYLE_ENTRY_ID,
                    "company": "甲公司",
                    "role": "工程师",
                    "start_date": "2020-01",
                    "end_date": "2021-01",
                    "description": ["做了一件事"],
                }
            ],
            "campus_experience": [],
            "projects": [],
            "skills": [],
            "awards": [],
        }
    )


def _styled_document(config: dict | None):
    result = build_resume_docx(
        _styled_resume(), template="classic", page_limit=2, format_config=config
    )
    return Document(BytesIO(result.content))


def _bullet_paragraph(document):
    return next(p for p in document.paragraphs if "做了一件事" in p.text)


def test_docx_body_indent_uses_the_object_role_override():
    """「正文 → 缩进」只动要点缩进，不动标题/日期（Q12 正文范围写 object_roles）。"""
    from app.services.pdf_exporter import resolve_layout
    from app.services.resume.resume_templates import FONT_SCALES

    base = float(FONT_SCALES["standard"]["base_px"])
    layout = resolve_layout(template="classic", base_px=base, format_config=None)
    plain = _bullet_paragraph(_styled_document(None))
    styled = _bullet_paragraph(
        _styled_document(
            {
                "object_roles": {
                    "version": 1,
                    f"{_STYLE_ENTRY_ID}:body": {"indent": 1.2},
                }
            }
        )
    )
    expected = Mm(layout.gap_mm(1.2)).twips
    assert styled.paragraph_format.left_indent.twips == pytest.approx(expected, abs=1)
    assert styled.paragraph_format.first_line_indent.twips == pytest.approx(-expected, abs=1)
    # 没覆盖时仍按模板 li_indent（Q11 保持既有口径）。
    assert plain.paragraph_format.left_indent.twips == pytest.approx(
        Mm(layout.gap_mm(layout.li_indent)).twips, abs=1
    )


def test_docx_body_align_and_bold_italic_reach_the_bullet_paragraph():
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    document = _styled_document(
        {
            "object_roles": {
                "version": 1,
                f"{_STYLE_ENTRY_ID}:body": {"align": "center", "bold": True, "italic": True},
            }
        }
    )
    paragraph = _bullet_paragraph(document)
    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.runs[-1]
    assert run.font.bold is True
    assert run.font.italic is True


def test_docx_letter_spacing_lands_as_ooxml_spacing_before_sz():
    """字距 → w:spacing（1/20 磅），且必须插在 w:sz 之前（rPr 子元素顺序由 schema 固定）。"""
    from app.services.resume.resume_templates import FONT_SCALES

    base = float(FONT_SCALES["standard"]["base_px"])
    document = _styled_document(
        {
            "object_roles": {
                "version": 1,
                f"{_STYLE_ENTRY_ID}:body": {"letter_spacing": 0.05},
            }
        }
    )
    run = _bullet_paragraph(document).runs[-1]
    rpr = run._element.find(qn("w:rPr"))
    spacing = rpr.find(qn("w:spacing"))
    assert spacing is not None, "字距必须落到 w:spacing"
    # 1 px = 0.75 pt；w:spacing 的单位是 1/20 磅。
    assert spacing.get(qn("w:val")) == str(int(round(base * 0.05 * 0.75 * 20)))
    tags = [child.tag.split("}")[-1] for child in rpr]
    assert tags.index("spacing") < tags.index("sz")


def test_docx_meta_align_inline_keeps_the_date_next_to_the_title():
    """「灰字 → 日期位置」在 Word 里 = 不用右制表位、日期紧随标题（Q12 灰字范围）。"""
    document = _styled_document(
        {
            "object_roles": {
                "version": 1,
                f"{_STYLE_ENTRY_ID}:meta": {"meta_align": "inline"},
            }
        }
    )
    head = next(p for p in document.paragraphs if "甲公司" in p.text)
    assert "·" in head.text
    assert "<w:tab/>" not in head._p.xml

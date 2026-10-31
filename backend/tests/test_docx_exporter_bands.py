# SPDX-License-Identifier: Apache-2.0
"""Word 色块：模板的页面底纹 / 实心标题块在 .docx 里的映射（w:shd + w:pBdr）。

从 test_docx_exporter.py 拆出（2026-10-31，纯搬迁）：原文件超过 500 行预算；
本组用例只守"Word 里的色块"这一块模板映射，自带最小夹具。
"""

from io import BytesIO

from app.schemas.resume import ResumeContent
from app.services.docx_exporter import build_resume_docx
from docx import Document
from docx.oxml.ns import qn

# ── Word 里的色块（2026-10-24 用户口径："导出的 word 也有绿色背景"）──────────
#
# Word 没有 CSS 的 `background`：能表达"色块"的只有**段落底纹**（`w:shd`）与
# **段落边框**（`w:pBdr`）。这里逐模板钉住映射关系——映射错了不会报错，只会让
# 用户导出的 Word 又变回"白底黑字"，与预览对不上。


def _shading_fills(document) -> list[str]:
    """文档里所有段落底纹的填充色（`RRGGBB`）。"""
    return [node.get(qn("w:fill")) for node in document.element.body.findall(".//" + qn("w:shd"))]


def _header_resume() -> ResumeContent:
    return ResumeContent(
        name="张三",
        phone="13800000000",
        email="zhangsan@example.com",
        job_intent="数据分析",
        summary="一段总结。",
    )


def test_modern_word_gets_the_soft_green_bands():
    """modern：页头三段 + 区块标题都要有浅绿底纹（`--accent-soft`）。"""
    from app.services.resume.resume_templates import TEMPLATE_LAYOUT_DEFAULTS

    expected = str(TEMPLATE_LAYOUT_DEFAULTS["modern"]["accent_soft"]).lstrip("#").upper()
    document = Document(BytesIO(build_resume_docx(_header_resume(), template="modern").content))

    fills = _shading_fills(document)
    # 页头：姓名 / 意向 / 联系方式 三段；再加「个人总结」这个区块标题。
    assert fills.count(expected) >= 4
    # 底纹必须配一圈同色边框把色带撑开（模拟 CSS 的 padding），否则色带贴着字。
    assert document.element.body.xml.count("<w:pBdr") >= 4


def test_technical_word_section_title_is_filled_with_the_accent_and_white_text():
    """technical：实心强调色块 + 白字（`section_title_style = accent_box`）。"""
    from app.services.resume.resume_labels import section_titles
    from app.services.resume.resume_templates import TEMPLATE_LAYOUT_DEFAULTS

    accent = str(TEMPLATE_LAYOUT_DEFAULTS["technical"]["accent"]).lstrip("#").upper()
    document = Document(BytesIO(build_resume_docx(_header_resume(), template="technical").content))

    assert _shading_fills(document) == [accent]
    summary_title = section_titles("zh-CN")["summary"]
    title = next(p for p in document.paragraphs if p.text == summary_title)
    colors = [run.font.color.rgb for run in title.runs if run.font.color and run.font.color.rgb]
    assert [str(color) for color in colors] == ["FFFFFF"]


def test_classic_word_section_title_has_a_left_bar_and_no_fill():
    """classic：只有左侧竖条，没有底纹（`section_title_style = left_bar`）。"""
    document = Document(BytesIO(build_resume_docx(_header_resume(), template="classic").content))

    assert _shading_fills(document) == []
    xml = document.element.body.xml
    assert '<w:left w:val="single"' in xml
    # 竖条是"只画左边"：不连底边一起画，否则会变成一条下划线。
    assert "<w:bottom w:val" not in xml

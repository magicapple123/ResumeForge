# SPDX-License-Identifier: Apache-2.0
"""Word 出口（v2）：可打开/文本可提取、条目头右制表位、lead 无框表格、
start_on_new_page 分页符、stacked 条目头、隐藏项不出现。"""

from __future__ import annotations

import copy

from app.services.docx_exporter import build_resume_docx
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from tests.document_export_samples import (
    HIDDEN_VALUES,
    ORDERED_VALUES,
    docx_text,
    make_document,
    normalize_text,
)


def _build(document=None, **kwargs):
    return build_resume_docx(document if document is not None else make_document(), **kwargs)


def test_docx_is_readable_and_contains_fields():
    result = _build()
    text = docx_text(result.content)
    # 质量门：文件能被 python-docx 打开（= 合法 OOXML），文本可提取。
    Document(__import__("io").BytesIO(result.content))
    for value in ORDERED_VALUES:
        assert value in text, f"Word 缺少字段值：{value}"
    positions = [text.index(value) for value in ORDERED_VALUES]
    assert positions == sorted(positions)
    for hidden in HIDDEN_VALUES:
        assert hidden not in text
    assert result.pages is None or result.pages >= 1


def test_docx_entry_head_uses_right_tab_stop():
    result = _build()
    document = Document(__import__("io").BytesIO(result.content))
    head_paragraphs = [p for p in document.paragraphs if "\t" in p.text]
    assert head_paragraphs, "条目头应使用制表位把元信息推到版心右缘"
    head = head_paragraphs[0]
    assert "示例大学" in head.text and "本科" in head.text
    stops = list(head.paragraph_format.tab_stops)
    assert stops and any(stop.alignment == WD_TAB_ALIGNMENT.RIGHT for stop in stops)


def test_docx_equal_lead_rows_use_borderless_table():
    result = _build()
    document = Document(__import__("io").BytesIO(result.content))
    assert document.tables, "lead/trail 的 equal 布局应落成无框表格"
    cells = [
        normalize_text(cell.text)
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    joined = " ".join(cells)
    assert "3.7/4.0" in joined
    assert "市场调研" in joined


def test_docx_stacked_head_puts_meta_on_next_paragraph():
    result = _build()
    document = Document(__import__("io").BytesIO(result.content))
    texts = [normalize_text(p.text) for p in document.paragraphs]
    title_index = next(index for index, text in enumerate(texts) if "校园招聘会策划" in text)
    # stacked：标题一行、元信息下一行（同一段里不再左右分栏）。
    assert "项目负责人" not in texts[title_index]
    assert "项目负责人" in texts[title_index + 1]


def test_docx_start_on_new_page_emits_page_break():
    document = copy.deepcopy(make_document())
    for section in document["sections"]:
        if section["key"] == "projects":
            section["layout"] = {"start_on_new_page": True}
    result = _build(document)
    from tests.document_export_samples import docx_page_break_count

    assert docx_page_break_count(result.content) >= 1


def test_docx_without_photo_has_no_inline_shapes():
    document = copy.deepcopy(make_document())
    document["header"]["items"] = [
        item for item in document["header"]["items"] if item.get("placement") != "photo"
    ]
    result = _build(document)
    doc = Document(__import__("io").BytesIO(result.content))
    assert len(doc.inline_shapes) == 0

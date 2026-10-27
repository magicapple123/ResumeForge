# SPDX-License-Identifier: Apache-2.0
"""Word（.docx）导出：字节可读、版式复用 ResumeLayout、分页与 PDF ±1。"""

from io import BytesIO

import pytest
from app.schemas.resume import ResumeContent
from app.services.docx_exporter import build_resume_docx
from app.services.pdf_exporter import (
    PAGE_WIDTH_MM,
    build_resume_pdf,
    font_available,
    resolve_layout,
)
from app.services.resume.resume_sample import sample_resume_content
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Mm

needs_font = pytest.mark.skipif(not font_available(), reason="本机没有可用的中文字体")

# 一张合法的 1x1 PNG（data URL）：测"照片进没进文档"不需要真实尺寸。
PHOTO = (
    "data:image/png;base64,"
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "YAAAAAYAAjCB0C8AAAAASUVORK5CYII="
)


def anchors(document):
    """文档里所有**浮动**图片（`wp:anchor`）。

    `python-docx` 的 `document.inline_shapes` 只认嵌入型（`wp:inline`），照片改成
    「浮于文字上方」之后它一个都看不到——断言必须走原始 XML，否则会得到"图片丢了"的假结论。
    """
    return document.element.body.findall(".//" + qn("wp:anchor"))


def _resume() -> ResumeContent:
    return ResumeContent(
        name="张三",
        phone="13800000000",
        summary="负责后端服务的设计与开发。",
        experience=[{"company": "字节跳动", "role": "工程师", "description": ["实现检索接口"]}],
        skills=[{"name": "Python", "level": "熟练"}],
    )


def test_docx_builds_a_valid_word_document():
    result = build_resume_docx(_resume(), template="classic", page_limit=1)

    # .docx 是 ZIP 容器，文件头以 PK 开头；且能被 python-docx 重新打开。
    assert result.content.startswith(b"PK")
    document = Document(BytesIO(result.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "张三" in text
    assert "字节跳动" in text
    assert "实现检索接口" in text


def test_docx_reuses_layout_margins_and_respects_margin_override():
    result = build_resume_docx(_resume(), template="classic", page_limit=1, margin_mm=20)

    document = Document(BytesIO(result.content))
    section = document.sections[0]
    # 旧的兼容输入（只给 margin_mm）仍按它排版，且四边一致
    # （复用 ResumeLayout.margin_mm 的同一来源）。
    assert round(section.top_margin.mm, 2) == pytest.approx(20.0, abs=0.01)
    assert round(section.left_margin.mm, 2) == pytest.approx(20.0, abs=0.01)
    assert round(section.right_margin.mm, 2) == pytest.approx(20.0, abs=0.01)
    assert round(section.bottom_margin.mm, 2) == pytest.approx(20.0, abs=0.01)


def test_docx_page_padding_beats_the_legacy_margin_mm():
    """Word 与 PDF 同一条规则（R1）：``page_padding`` 权威，``margin_mm`` 不再覆盖。

    修之前 ``margin_mm=30`` 会盖掉 ``page_padding=8``，于是同一份导出里 Word/PDF 是
    30mm 而预览/HTML 是 8mm——"改面板值预览不变、导出变"的同一个根因。
    """
    result = build_resume_docx(
        _resume(),
        template="classic",
        page_limit=1,
        format_config={"page_padding": 8},
        margin_mm=30,
    )

    section = Document(BytesIO(result.content)).sections[0]
    assert round(section.left_margin.mm, 2) == pytest.approx(8.0, abs=0.01)
    assert round(section.top_margin.mm, 2) == pytest.approx(8.0, abs=0.01)


def test_docx_clamps_out_of_range_legacy_margin_mm():
    """越界的兼容值收进 0–26（不 500、不把正文排出纸张）。"""
    result = build_resume_docx(_resume(), template="classic", page_limit=1, margin_mm=40)

    section = Document(BytesIO(result.content)).sections[0]
    assert round(section.left_margin.mm, 2) == pytest.approx(26.0, abs=0.01)


def test_docx_uses_template_default_margin_when_not_overridden():
    from app.services.resume.resume_templates import TEMPLATE_LAYOUT_DEFAULTS

    expected = float(TEMPLATE_LAYOUT_DEFAULTS["classic"]["padding_mm"])
    result = build_resume_docx(_resume(), template="classic", page_limit=1)

    document = Document(BytesIO(result.content))
    section = document.sections[0]
    assert round(section.top_margin.mm, 2) == pytest.approx(expected, abs=0.01)


def test_docx_page_is_a4_not_the_python_docx_default_letter():
    """纸张必须是 A4（210×297mm），与预览和 PDF 同一张纸。

    python-docx 默认是 Letter（215.9×279.4mm）：不显式设置时，同一段文字在 Word 里
    行宽与页高都不同，"Word 页数 ≈ PDF 页数"这个承诺从纸型上就不成立。
    """
    result = build_resume_docx(_resume(), template="classic", page_limit=1)

    section = Document(BytesIO(result.content)).sections[0]
    assert round(section.page_width.mm, 2) == pytest.approx(210.0, abs=0.05)
    assert round(section.page_height.mm, 2) == pytest.approx(297.0, abs=0.05)


def test_docx_section_margins_follow_all_four_sides():
    """四向页边距 1:1 落进 ``w:pgMar``：``Mm()`` 与既有路径同一个构造器，无单位坑。

    ``(上, 右, 下, 左)`` 里只给出三条时，缺的边跟"总"走——与 PDF 的解析口径逐字相同
    （``layout._resolve_page_margins``）。Word 与 PDF 若各算一遍，同一次导出就会给出
    两种边距，而这在界面上完全看不出来（用户只看 PDF）。
    """
    result = build_resume_docx(
        _resume(),
        template="classic",
        page_limit=1,
        format_config={
            "page_padding": 14,
            "page_padding_top": 8,
            "page_padding_bottom": 12,
        },
    )

    section = Document(BytesIO(result.content)).sections[0]
    assert round(section.top_margin.mm, 2) == pytest.approx(8.0, abs=0.01)
    assert round(section.bottom_margin.mm, 2) == pytest.approx(12.0, abs=0.01)
    # 没给的两条边跟"总"（14）走。
    assert round(section.left_margin.mm, 2) == pytest.approx(14.0, abs=0.01)
    assert round(section.right_margin.mm, 2) == pytest.approx(14.0, abs=0.01)


def test_docx_four_equal_sides_stay_equal_to_the_legacy_value():
    """四边相等时与"只有一个 margin_mm"的历史行为逐点相同（存量文档不变形）。"""
    uniform = build_resume_docx(_resume(), template="classic", page_limit=1, margin_mm=20)
    explicit = build_resume_docx(
        _resume(), template="classic", page_limit=1, format_config={"page_padding": 20}
    )

    for build in (uniform, explicit):
        section = Document(BytesIO(build.content)).sections[0]
        for margin in (
            section.top_margin,
            section.right_margin,
            section.bottom_margin,
            section.left_margin,
        ):
            assert round(margin.mm, 2) == pytest.approx(20.0, abs=0.01)


def test_docx_excludes_photo_when_include_photo_is_false():
    # 用一张合法 PNG 的 data URL；关闭照片后文档里不应出现任何图片。
    png = (
        "data:image/png;base64,"
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
        "YAAAAAYAAjCB0C8AAAAASUVORK5CYII="
    )
    resume = ResumeContent(name="张三", photo=png, summary="有照片。")

    with_photo = build_resume_docx(resume, template="classic", page_limit=1, include_photo=True)
    without_photo = build_resume_docx(resume, template="classic", page_limit=1, include_photo=False)

    # 带照片的文档里有且只有 1 张图片，关掉后一张也没有。
    # 图片是**浮动**的（`wp:anchor`），所以 `inline_shapes` 会是 0——见下面那条用例。
    doc_with = Document(BytesIO(with_photo.content))
    doc_without = Document(BytesIO(without_photo.content))
    assert len(anchors(doc_with)) == 1
    assert len(anchors(doc_without)) == 0


def test_docx_gender_is_a_separate_muted_run_not_the_name_run():
    """Word 里性别必须是独立的 muted 灰 run，不再与姓名拼成一个强调色 run。

    此前 `name_line = name + "　" + gender` 拼成一个 run，性别跟着姓名一起变大变粗、
    变强调色。拆分后姓名 run 是粗体 + 强调色 + name_ratio 字号，性别 run 是常规 + muted +
    name_extra_ratio 字号。
    """
    from app.services.resume.resume_templates import TEMPLATE_LAYOUT_DEFAULTS

    resume = ResumeContent(name="张三", gender="男", summary="一句话总结。")
    result = build_resume_docx(resume, template="classic", page_limit=1)
    doc = Document(BytesIO(result.content))

    # 第一个段落应是页头照片段或姓名段；找到含"张三"的段落。
    name_paragraph = next(p for p in doc.paragraphs if "张三" in p.text)
    runs = name_paragraph.runs
    assert len(runs) >= 2, "姓名与性别应分成两个 run"
    name_run = next(r for r in runs if "张三" in r.text)
    gender_run = next(r for r in runs if "男" in r.text)

    defaults = TEMPLATE_LAYOUT_DEFAULTS["classic"]
    base_px = 14.0
    # 姓名：粗体、强调色、name_ratio 字号
    assert name_run.font.bold is True
    # python-docx 把 Pt 量化到 0.5pt（19.53 → 19.5），用 0.05 容差。
    assert name_run.font.size.pt == pytest.approx(
        base_px * float(defaults["name_ratio"]) * 0.75, abs=0.05
    )
    # 性别：非常规粗体（None 或 False）、name_extra_ratio 字号、且字号小于姓名
    assert not gender_run.font.bold
    assert gender_run.font.size.pt == pytest.approx(
        base_px * float(defaults["name_extra_ratio"]) * 0.75, abs=0.05
    )
    assert gender_run.font.size.pt < name_run.font.size.pt


@needs_font
def test_docx_page_estimate_is_within_one_of_pdf():
    resume = sample_resume_content()
    pdf = build_resume_pdf(resume, template="classic", page_limit=2, font_scale="standard")
    docx = build_resume_docx(resume, template="classic", page_limit=2, font_scale="standard")

    assert docx.pages is not None
    assert abs(docx.pages - pdf.pages) <= 1


# 与预览对齐的两处硬约束（2026-10-24 用户口径："Word 版和预览差别很大"）。
#
# 1. 条目头的日期必须**右对齐**：预览用 `justify-content: space-between`，Word 用
#    右对齐制表位。此前是"标题 + 全角空格 + 日期"，日期跟着标题长度浮动。
# 2. 项目符号必须**悬挂缩进**：折行的第二行要回到正文起点，而不是退回版心左边缘。


def test_entry_head_uses_a_right_aligned_tab_stop():
    document = Document(BytesIO(build_resume_docx(_resume(), page_limit=1).content))
    heads = [
        paragraph
        for paragraph in document.paragraphs
        if len(paragraph.paragraph_format.tab_stops) > 0
    ]
    assert heads, "条目头没有右对齐制表位，日期会跟着标题长度左右浮动"
    for paragraph in heads:
        stop = paragraph.paragraph_format.tab_stops[0]
        assert stop.alignment == WD_TAB_ALIGNMENT.RIGHT
        # 制表位必须落在版心右边缘，而不是某个随手写的固定位置。
        layout = resolve_layout(template="classic", base_px=15.0, format_config=None)
        expected_twips = Mm(PAGE_WIDTH_MM - layout.margin_left - layout.margin_right).twips
        # Word 内部以 twip 存位置，取整误差允许 1 twip（≈0.018mm）。
        assert abs(stop.position.twips - expected_twips) <= 1


def test_bullets_hang_so_wrapped_lines_align_with_the_text():
    """悬挂缩进必须等于模板 li_indent 换算出的毫米（W07 修掉了写死 3.0mm）。

    W07 之前 Word 固定 3.0mm、PDF/HTML 读模板值（classic 1.43 × 字号 ≈ 5.3mm），三种格式的
    列表缩进不一致；现在 Word 也从 layout.li_indent 派生——**这是有意的可见变化**。
    """
    from app.services.pdf_exporter import resolve_layout
    from app.services.resume.resume_templates import FONT_SCALES

    layout = resolve_layout(
        template="classic", base_px=float(FONT_SCALES["standard"]["base_px"]), format_config=None
    )
    expected_mm = layout.gap_mm(layout.li_indent)
    assert expected_mm != 3.0, "classic 的模板缩进不是 3.0mm，这条断言才有效"

    document = Document(BytesIO(build_resume_docx(_resume(), page_limit=1).content))
    bullets = [p for p in document.paragraphs if p.text.startswith("• ")]
    assert bullets, "列表项没有用项目符号"
    for paragraph in bullets:
        # twip 取整：毫米换算成 twip 后允许 1 twip 误差。
        assert abs(paragraph.paragraph_format.first_line_indent.twips + Mm(expected_mm).twips) <= 1
        assert abs(paragraph.paragraph_format.left_indent.twips - Mm(expected_mm).twips) <= 1


# ── 照片必须是「浮于文字上方」（2026-10-24 用户实测的解法）──────────────────
#
# 嵌入型图片是正文流里的一个"大号字符"：独占一行、把姓名整体下推，页头因此比预览高
# 一截，后面的分页也跟着错。改成浮动之后不占行高，页头与预览同高。
# 「浮于文字上方」= `wp:anchor` + `behindDoc="0"` + `<wp:wrapNone/>`；
# 「衬于文字下方」是同一个 `wrapNone`、`behindDoc="1"`——这一条必须钉死，写反了照片
# 会跑到文字后面，比嵌入型更糟。


def test_photo_floats_in_front_of_text_instead_of_occupying_a_line():
    resume = ResumeContent(name="张三", photo=PHOTO, summary="有照片。")
    document = Document(BytesIO(build_resume_docx(resume, page_limit=1).content))

    anchor = anchors(document)[0]
    assert anchor.get("behindDoc") == "0"
    assert anchor.find(qn("wp:wrapNone")) is not None
    # 水平贴版心右缘（classic 的照片在右），垂直锚在段落上——翻页时不会跑到纸张顶端。
    position_h = anchor.find(qn("wp:positionH"))
    assert position_h.get("relativeFrom") == "margin"
    assert position_h.find(qn("wp:align")).text == "right"
    assert anchor.find(qn("wp:positionV")).get("relativeFrom") == "paragraph"

    # 照片挂在**姓名那一段**里，不再单独占一行（这正是"页头矮了一截"的来源）。
    name_paragraph = next(p for p in document.paragraphs if p.text.startswith("张三"))
    assert name_paragraph._p.find(".//" + qn("w:drawing")) is not None


def test_photo_float_side_follows_the_template():
    """elegant 的 `.header` 是 `row-reverse`（照片在左），浮动锚点必须跟着贴左。"""
    resume = ResumeContent(name="张三", photo=PHOTO, summary="有照片。")
    document = Document(BytesIO(build_resume_docx(resume, template="elegant").content))

    position_h = anchors(document)[0].find(qn("wp:positionH"))
    assert position_h.find(qn("wp:align")).text == "left"


# ── Word 里的色块（2026-10-24 用户口径："导出的 word 也有绿色背景"）──────────
#
# Word 没有 CSS 的 `background`：能表达"色块"的只有**段落底纹**（`w:shd`）与
# **段落边框**（`w:pBdr`）。这里逐模板钉住映射关系——映射错了不会报错，只会让
# 用户导出的 Word 又变回"白底黑字"，与预览对不上。


def _shading_fills(document) -> list[str]:
    """文档里所有段落底纹的填充色（`RRGGBB`）。"""
    return [node.get(qn("w:fill")) for node in document.element.body.findall(".//" + qn("w:shd"))]


def _header_resume() -> ResumeContent:
    return ResumeContent(
        name="张三",
        phone="13800000000",
        email="zhangsan@example.com",
        job_intent="数据分析",
        summary="一段总结。",
    )


def test_modern_word_gets_the_soft_green_bands():
    """modern：页头三段 + 区块标题都要有浅绿底纹（`--accent-soft`）。"""
    from app.services.resume.resume_templates import TEMPLATE_LAYOUT_DEFAULTS

    expected = str(TEMPLATE_LAYOUT_DEFAULTS["modern"]["accent_soft"]).lstrip("#").upper()
    document = Document(BytesIO(build_resume_docx(_header_resume(), template="modern").content))

    fills = _shading_fills(document)
    # 页头：姓名 / 意向 / 联系方式 三段；再加「个人总结」这个区块标题。
    assert fills.count(expected) >= 4
    # 底纹必须配一圈同色边框把色带撑开（模拟 CSS 的 padding），否则色带贴着字。
    assert document.element.body.xml.count("<w:pBdr") >= 4


def test_technical_word_section_title_is_filled_with_the_accent_and_white_text():
    """technical：实心强调色块 + 白字（`section_title_style = accent_box`）。"""
    from app.services.resume.resume_labels import section_titles
    from app.services.resume.resume_templates import TEMPLATE_LAYOUT_DEFAULTS

    accent = str(TEMPLATE_LAYOUT_DEFAULTS["technical"]["accent"]).lstrip("#").upper()
    document = Document(BytesIO(build_resume_docx(_header_resume(), template="technical").content))

    assert _shading_fills(document) == [accent]
    summary_title = section_titles("zh-CN")["summary"]
    title = next(p for p in document.paragraphs if p.text == summary_title)
    colors = [run.font.color.rgb for run in title.runs if run.font.color and run.font.color.rgb]
    assert [str(color) for color in colors] == ["FFFFFF"]


def test_classic_word_section_title_has_a_left_bar_and_no_fill():
    """classic：只有左侧竖条，没有底纹（`section_title_style = left_bar`）。"""
    document = Document(BytesIO(build_resume_docx(_header_resume(), template="classic").content))

    assert _shading_fills(document) == []
    xml = document.element.body.xml
    assert '<w:left w:val="single"' in xml
    # 竖条是"只画左边"：不连底边一起画，否则会变成一条下划线。
    assert "<w:bottom w:val" not in xml


# ===== W08：孤行/寡行控制的 OOXML 落点 =====
#
# Word 没有 CSS 的 break-inside / break-after，对应的 OOXML 属性是 w:keepLines（整段不跨页）
# 与 w:keepNext（跟随后续内容）。这里直接看 pPr 的子元素名，而不是只看"文档能打开"——
# 少一个属性不会报错，只会在打印时把一条要点劈成两页。


def _pPr_flags(paragraph) -> set[str]:
    """段落属性里出现的 OOXML 标签名（不含命名空间）。"""
    pPr = paragraph._p.find(qn("w:pPr"))
    if pPr is None:
        return set()
    return {child.tag.split("}")[-1] for child in pPr}


def test_docx_lands_keep_flags_for_titles_and_paragraphs():
    from app.services.resume.resume_labels import section_titles

    result = build_resume_docx(_resume(), template="classic", page_limit=1)
    document = Document(BytesIO(result.content))
    paragraphs = document.paragraphs

    title_text = section_titles("zh-CN")["experience"]
    title = next(paragraph for paragraph in paragraphs if paragraph.text == title_text)
    assert "keepNext" in _pPr_flags(title), "板块标题必须跟随后续内容（w:keepNext）"

    entry_head = next(paragraph for paragraph in paragraphs if "字节跳动" in paragraph.text)
    assert "keepNext" in _pPr_flags(entry_head), "条目头必须跟着它的正文（w:keepNext）"

    bullet = next(paragraph for paragraph in paragraphs if "实现检索接口" in paragraph.text)
    assert "keepLines" in _pPr_flags(bullet), "一条要点不跨页（w:keepLines）"

    summary = next(paragraph for paragraph in paragraphs if "负责后端服务" in paragraph.text)
    assert "keepLines" in _pPr_flags(summary), "摘要段不跨页（w:keepLines）"


def test_docx_keep_flags_are_not_written_for_plain_header_paragraphs():
    """默认不写属性：页头的姓名段不该长出 keepNext/keepLines（旧产物不被动过）。"""
    result = build_resume_docx(_resume(), template="classic", page_limit=1)
    document = Document(BytesIO(result.content))
    name = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("张三"))

    flags = _pPr_flags(name)
    assert "keepNext" not in flags
    assert "keepLines" not in flags

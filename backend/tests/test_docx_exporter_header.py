# SPDX-License-Identifier: Apache-2.0
"""Word 页头回归：联系方式字段集合 / 强制换行 / 照片取景（#6、#3、#10）。

拆分自 `test_docx_exporter.py` 的 500 行测试预算（`test_file_size_budget.py`）：
页头这几条用例只关心 .docx 解包后的段落/图片几何，与"正文分区怎么排"不是同一件事。

从 `test_docx_exporter.py` 单向导入共享 helper（兄弟文件单向导入是既有约定）。
"""

from __future__ import annotations

from io import BytesIO

import pytest
from app.schemas.resume import ResumeContent
from app.services.docx_exporter import build_resume_docx
from docx import Document
from docx.oxml.ns import qn
from test_docx_exporter import anchors


def _body_text(document) -> str:
    """文档正文的纯文本（段落按换行拼接）。"""
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_docx_photo_keeps_its_own_aspect_ratio_inside_the_photo_box():
    """Word 的照片按**原图比例**装进模板照片框（object-fit: contain，2026-10 #10）。

    旧实现先按框比例 cover 裁剪，横图会被裁掉两侧、硬变成竖图；现在宽高比必须等于
    原图比例，且高度不超过照片框高度（框内留白，不溢出）。这是 `photo_crop` 生效的
    前提：服务端已派生出裁剪图，任何再裁一次都会让用户"看不到裁剪效果"。
    """
    import base64
    from io import BytesIO

    from app.services.resume.resume_templates import TEMPLATE_LAYOUT_DEFAULTS
    from PIL import Image

    # 用 Pillow 现造一张 4:1 的横图 PNG（与 classic 照片框 0.7875 比例明显不同）。
    buf = BytesIO()
    Image.new("RGB", (400, 100), (0, 128, 200)).save(buf, format="PNG")
    png = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode("ascii")
    resume = ResumeContent(name="张三", photo=png, summary="带照片。")
    result = build_resume_docx(resume, template="classic", page_limit=1, include_photo=True)
    doc = Document(BytesIO(result.content))
    extent = anchors(doc)[0].find(qn("wp:extent"))
    cx, cy = int(extent.get("cx")), int(extent.get("cy"))

    # 宽度应等于模板照片框宽度（mm 换算成 EMU）。
    defaults = TEMPLATE_LAYOUT_DEFAULTS["classic"]
    expected_width_emu = int(round(14.0 * float(defaults["photo_width_ratio"]) * 0.264583 * 36000))
    assert abs(cx - expected_width_emu) < (expected_width_emu * 0.1)
    # 宽高比 = 原图 4:1；旧 cover 实现会得到 6.3:8 ≈ 0.79（照片被裁成竖图）。
    assert cx / cy == pytest.approx(4.0, rel=1e-3)


def test_docx_photo_crop_params_decide_the_embedded_frame():
    """`photo_crop` 派生图的比例 = Word 里照片的实际比例（裁剪参数真的决定取景）。

    40×20 的横图，裁成左半块 → 20×20 正方形。若 Word 还用 cover 按框比例再裁，
    比例会回到 6.3:8；用 contain 则严格等于派生图比例 1:1。
    """
    import base64
    import hashlib
    from io import BytesIO

    from PIL import Image

    image = Image.new("RGB", (40, 20))
    image.putdata(
        [
            ((x * 6) % 256, (y * 12) % 256, (x * 3 + y * 7) % 256)
            for y in range(20)
            for x in range(40)
        ]
    )
    buf = BytesIO()
    image.save(buf, format="PNG")
    png = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode("ascii")
    crop = {
        "version": 1,
        "source": {
            "sha256": hashlib.sha256(png.encode("utf-8")).hexdigest(),
            "width": 40,
            "height": 20,
        },
        "rect": {"x": 0.0, "y": 0.0, "w": 0.5, "h": 1.0},
        "aspect": "free",
    }
    resume = ResumeContent(name="张三", photo=png, summary="带照片。")
    result = build_resume_docx(
        resume,
        template="classic",
        page_limit=1,
        include_photo=True,
        format_config={"photo_crop": crop},
    )
    extent = anchors(Document(BytesIO(result.content)))[0].find(qn("wp:extent"))
    cx, cy = int(extent.get("cx")), int(extent.get("cy"))
    assert cx / cy == pytest.approx(1.0, rel=1e-3), "Word 没有按 photo_crop 派生图的比例取景"


# ===== 页眉联系方式与预览同集合（2026-10 #6）=====
#
# 用户反馈"预览里有 GitHub 和个人主页，导出的 Word 里没有"：Word 页头只拼了
# phone / email / city / birth_year 四个字段。下面从**真实 .docx 解包后的文本**里核对，
# 而不是断言拼接用的那个字符串——后者只能证明"我拼对了"，不能证明"文档里真的有"。

CONTACT_RESUME = ResumeContent(
    name="张三",
    phone="13800001234",
    email="zhangsan@example.com",
    city="杭州",
    birth_year="1998",
    github="https://github.com/zhangsan",
    personal_website="https://zhangsan.dev",
    summary="一句话总结。",
)


def _body_text(document) -> str:
    """文档正文的纯文本（段落按换行拼接）。"""
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_docx_header_contact_keeps_every_field_the_preview_renders():
    """Word 页头必须包含预览里的**全部六个**联系方式字段（含标签文案）。"""
    document = Document(
        BytesIO(build_resume_docx(CONTACT_RESUME, template="classic", page_limit=1).content)
    )
    text = _body_text(document)

    for value in (
        "13800001234",
        "zhangsan@example.com",
        "杭州",
        "1998",
        "GitHub：https://github.com/zhangsan",
        "个人主页：https://zhangsan.dev",
    ):
        assert value in text, f"Word 页头缺少联系方式字段：{value!r}"


def test_docx_header_contact_forced_breaks_are_word_line_breaks():
    """`contact_breaks` 在 Word 里落成 `w:br`（与 HTML 的 `<br>` 同义），默认零个。

    Word 没有 CSS 的 `white-space`：强制换行只能写 OOXML `w:br`。默认（旧简历没有
    这个键）一个字节都不加——否则"没设过换行的简历"会平白多一行。
    """
    default_doc = Document(
        BytesIO(build_resume_docx(CONTACT_RESUME, template="classic", page_limit=1).content)
    )
    assert default_doc.element.body.xml.count("<w:br") == 0

    broken_doc = Document(
        BytesIO(
            build_resume_docx(
                CONTACT_RESUME,
                template="classic",
                page_limit=1,
                format_config={"contact_breaks": ["personal_website", "github"]},
            ).content
        )
    )
    # 两个字段前各一个强制换行（第一个字段 phone 之前的 break 无意义，被丢弃）。
    assert broken_doc.element.body.xml.count("<w:br") == 2

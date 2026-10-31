# SPDX-License-Identifier: Apache-2.0
"""Word 照片锚定：嵌入型 -> 「浮于文字上方」（不占行高、页头高度与预览一致）。

从 test_docx_exporter.py 拆出（2026-10-31，纯搬迁）：原文件超过 500 行预算；
本组用例只守照片锚定这一段 OOXML 结构，共用主文件的 PHOTO / _resume / anchors。
"""

from io import BytesIO

from app.schemas.resume import ResumeContent
from app.services.docx_exporter import build_resume_docx
from docx import Document
from docx.oxml.ns import qn
from test_docx_exporter import PHOTO, anchors

# ── 照片必须是「浮于文字上方」（2026-10-24 用户实测的解法）──────────────────
#
# 嵌入型图片是正文流里的一个"大号字符"：独占一行、把姓名整体下推，页头因此比预览高
# 一截，后面的分页也跟着错。改成浮动之后不占行高，页头与预览同高。
# 「浮于文字上方」= `wp:anchor` + `behindDoc="0"` + `<wp:wrapNone/>`；
# 「衬于文字下方」是同一个 `wrapNone`、`behindDoc="1"`——这一条必须钉死，写反了照片
# 会跑到文字后面，比嵌入型更糟。


def test_photo_floats_in_front_of_text_instead_of_occupying_a_line():
    resume = ResumeContent(name="张三", photo=PHOTO, summary="有照片。")
    document = Document(BytesIO(build_resume_docx(resume, page_limit=1).content))

    anchor = anchors(document)[0]
    assert anchor.get("behindDoc") == "0"
    assert anchor.find(qn("wp:wrapNone")) is not None
    # 水平贴版心右缘（classic 的照片在右），垂直锚在段落上——翻页时不会跑到纸张顶端。
    position_h = anchor.find(qn("wp:positionH"))
    assert position_h.get("relativeFrom") == "margin"
    assert position_h.find(qn("wp:align")).text == "right"
    assert anchor.find(qn("wp:positionV")).get("relativeFrom") == "paragraph"

    # 照片挂在**姓名那一段**里，不再单独占一行（这正是"页头矮了一截"的来源）。
    name_paragraph = next(p for p in document.paragraphs if p.text.startswith("张三"))
    assert name_paragraph._p.find(".//" + qn("w:drawing")) is not None


def test_photo_float_side_follows_the_template():
    """elegant 的 `.header` 是 `row-reverse`（照片在左），浮动锚点必须跟着贴左。"""
    resume = ResumeContent(name="张三", photo=PHOTO, summary="有照片。")
    document = Document(BytesIO(build_resume_docx(resume, template="elegant").content))

    position_h = anchors(document)[0].find(qn("wp:positionH"))
    assert position_h.find(qn("wp:align")).text == "left"

# SPDX-License-Identifier: Apache-2.0
"""Word（.docx）简历导出（v2）：DocumentPlan -> python-docx。

**为什么不另写一套排版**：直出 PDF 与预览不一致的教训已经反复出现——任何渲染器只要
自己再存一份"行距 / 页边距 / 字号层级 / 强调色"，模板一改就会悄悄漂移。所以 Word 的
页边距、字号层级、行距、区块间距与强调色**全部**从 ResumeLayout 与 resume_templates 读
（与 PDF 同一来源）；内容结构（字段位置/顺序/条目头布局/列布局）只来自 DocumentPlan。

分页（Word 没有 fpdf 那种可回读的页数）用与 PDF **同一套连续流模型**估算：
连续内容高 ÷ 每页可用高。这个估算是"连续流"模型，与 Word 的自动分页一致，因此与
PDF 的真实页数差在 ±1 内（容差断言见 test_docx_exporter.py）。**不缩字号**：按版式参数
原样写入，估算页数超过 page_limit 时标记 overflow，由界面提示用户——Word 自己会顺延出
更多页，文档里不会丢内容。没有可用的中文字体时**仍能生成**（python-docx 只写字体名，
渲染由 Word 完成），但无法做页数估算，此时 pages=None。

入口接受 v1 ResumeContent（读时转换）或 v2 文档/JSON；converter 是唯一的版本归一化点，
四个出口看到的是同一套字段与顺序。
"""

from __future__ import annotations

import logging
import math
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.shared import Mm

from ..schemas.resume_generation import LONG_CV_MAX_PAGES
from .docx_blocks import (
    DOCX_FONT,
    DocxTheme,
    resolve_docx_theme,
)
from .docx_document_render import draw_document
from .pdf_exporter import (
    PAGE_HEIGHT_MM,
    PAGE_WIDTH_MM,
    _resolve_font_paths,
    measure_content_height,
    resolve_accent,
    resolve_accent_soft,
    resolve_font_paths,
    resolve_layout,
)
from .resume.document_html import (
    coerce_document,
    document_photo_value,
    paper_document,
    with_document_photo,
)
from .resume.document_plan import build_document_plan
from .resume.resume_labels import DEFAULT_RESUME_LANGUAGE, normalize_language
from .resume.resume_render_overrides import resolve_render_overrides
from .resume.resume_template_format import FONT_KEY
from .resume.resume_templates import (
    FONT_SCALES,
    template_spec,
    validated_format_config,
)

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class ResumeDOCX:
    """生成好的 Word 文档，连同估算页数与"装不装得下"的结论（语义与 ResumePDF 对齐）。"""

    content: bytes
    pages: int | None
    overflow: bool = False


def _role_style_query(overrides):
    """把 RenderOverrides 收成绘制层的查询函数；None = 模板默认（不做任何覆盖）。"""
    style = overrides.style if overrides is not None else None
    if style is None:
        return None

    def query(role: str, object_id: str | None = None):
        return style.effective(role, object_id=object_id)

    return query


def _set_run_font(run, *, size_pt: float, color=None) -> None:
    """续页信息（页眉/页脚）用的小字号 run：字体族与主题一致。"""
    from .docx_blocks import set_run_font

    set_run_font(run, size_pt=size_pt, color=color)


def _apply_continuation(doc, overrides, name: str) -> None:
    """W08：Word 的续页信息（重复姓名 / 页码）。

    画在 section.header / footer 的页边距区域里，不进正文流；默认 plan 全空 → 一个字节
    都不加（旧产物不变）。
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if overrides is None:
        return
    plan = overrides.pagination
    if plan.empty:
        return
    section = doc.sections[0]

    def _field(paragraph, instruction: str) -> None:
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = instruction
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(instr)
        run._r.append(end)

    if plan.repeat_name and name.strip():
        section.different_first_page_header_footer = True
        header_paragraph = section.header.paragraphs[0]
        header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header_paragraph.add_run(name)
        _set_run_font(run, size_pt=9)
    if plan.page_numbers:
        footer_paragraph = section.footer.paragraphs[0]
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        prefix = footer_paragraph.add_run("第 ")
        _set_run_font(prefix, size_pt=9)
        _field(footer_paragraph, "PAGE")
        middle = footer_paragraph.add_run(" 页 / 共 ")
        _set_run_font(middle, size_pt=9)
        _field(footer_paragraph, "NUMPAGES")
        suffix = footer_paragraph.add_run(" 页")
        _set_run_font(suffix, size_pt=9)


def _build_document(
    plan,
    *,
    accent: tuple[int, int, int],
    accent_soft: tuple[int, int, int],
    layout,
    include_photo: bool,
    theme: DocxTheme,
    overrides=None,
) -> bytes:
    """把 plan 写成 docx 字节：页面设置（A4 + 四边距）→ 内容 → 续页信息。"""
    doc = Document()
    section = doc.sections[0]
    # python-docx 默认是 Letter（215.9×279.4mm），而预览与 PDF 都是 A4（210×297mm）。
    # 不显式设置的话，同一段文字在 Word 里行宽与页高都不同，"页数与 PDF 一致"无从谈起。
    section.page_width = Mm(PAGE_WIDTH_MM)
    section.page_height = Mm(PAGE_HEIGHT_MM)
    section.top_margin = Mm(layout.margin_top)
    section.bottom_margin = Mm(layout.margin_bottom)
    section.left_margin = Mm(layout.margin_left)
    section.right_margin = Mm(layout.margin_right)

    draw_document(
        doc,
        plan,
        layout=layout,
        base=layout.base_px,
        include_photo=include_photo,
        theme=theme,
        accent=accent,
        accent_soft=accent_soft,
        role_style=_role_style_query(overrides),
    )
    name = "".join(str(seg.text or "") for seg in plan.header.name).strip()
    _apply_continuation(doc, overrides, name)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _estimate_pages(
    document,
    *,
    accent: tuple[int, int, int],
    layout,
    page_limit: int,
    include_photo: bool,
    language: str = DEFAULT_RESUME_LANGUAGE,
    font_key: str | None = None,
    overrides=None,
    format_config: dict | None = None,
) -> tuple[int | None, bool]:
    """估算 Word 页数与"是否超过所选页数"。

    Word 自动分页是"连续流"，所以页数用连续内容高 ÷ 每页可用高估算——与 PDF 的真实分页
    共用同一份版式与同一段测量代码（measure_content_height）。没有中文字体时返回
    (None, False)：页数未知不等于装得下，也不该凭空报溢出。
    """
    fonts = resolve_font_paths(font_key) if font_key is not None else _resolve_font_paths()
    if fonts is None:
        return None, False
    regular, bold = fonts
    height = measure_content_height(
        document,
        accent=accent,
        layout=layout,
        regular=regular,
        bold=bold,
        include_photo=include_photo,
        language=language,
        overrides=overrides,
        format_config=format_config,
    )
    pages = max(1, math.ceil(height / layout.usable_height_per_page))
    return pages, pages > page_limit


def build_resume_docx(
    content,
    *,
    template: str = "classic",
    page_limit: int = 1,
    font_scale: str = "standard",
    format_config: dict | None = None,
    margin_mm: float | None = None,
    include_photo: bool = True,
    language: str = DEFAULT_RESUME_LANGUAGE,
) -> ResumeDOCX:
    """把简历内容渲染成 Word 文档（入口接受 v1 或 v2 内容）。

    版式全部来自 ResumeLayout（复用 PDF 的同一来源），页数用连续流模型估算，与 PDF 的
    真实页数差在 ±1 内。没有中文字体时仍能生成文档，只是 pages=None。

    **不缩字号**：估算页数超过 page_limit 时只标记 overflow，文档本身按用户设定的版式
    写入——Word 会自己顺延页数，不会丢内容（与 PDF 的截断不同）。
    """
    resolved_language = normalize_language(language)
    # 投递视角（W06 / v2 hidden）：隐藏项在排版前剔除；页数估算与文档正文因此同源。
    document = paper_document(
        coerce_document(
            content, language=resolved_language, format_config=format_config
        )
    )
    spec = template_spec(template)
    scale = FONT_SCALES.get(font_scale) or FONT_SCALES["standard"]
    base_px = float(scale["base_px"])

    overrides = validated_format_config(format_config)
    adjust = overrides.get("font_scale_adjust")
    if isinstance(adjust, (int, float)):
        base_px = round(base_px * float(adjust), 2)
    accent = resolve_accent(spec["name"], overrides)
    accent_soft = resolve_accent_soft(spec["name"], overrides)
    # W05：Word 的字体族名（读所选字体的真实族名）与正文/辅助色（与 PDF 同源）。
    theme = resolve_docx_theme(spec["name"], overrides)
    # W07/W08/W15：角色/对象/分页/照片裁剪四域的唯一解析结果（纯读，不写库）。
    render_overrides = resolve_render_overrides(
        template=spec["name"],
        format_config=format_config,
        photo=document_photo_value(document),
    )
    if render_overrides.photo_crop is not None:
        document = with_document_photo(
            document, render_overrides.photo_for_render(document_photo_value(document))
        )

    layout = resolve_layout(
        template=spec["name"], base_px=base_px, format_config=overrides, margin_mm=margin_mm
    )
    # A4 长 CV：Word 只按 page_limit 估算 overflow、不截断内容，所以防御性上限用长 CV 的
    # 10 页（普通简历的 1–3 在调用方已按文档模式夹紧）。
    limit = max(1, min(int(page_limit), LONG_CV_MAX_PAGES))
    plan = build_document_plan(document, format_config, resolved_language)

    pages, overflow = _estimate_pages(
        document,
        accent=accent,
        layout=layout,
        page_limit=limit,
        include_photo=include_photo,
        language=resolved_language,
        font_key=overrides.get(FONT_KEY),
        overrides=render_overrides,
        format_config=overrides,
    )
    content = _build_document(
        plan,
        accent=accent,
        accent_soft=accent_soft,
        layout=layout,
        include_photo=include_photo,
        theme=theme,
        overrides=render_overrides,
    )
    return ResumeDOCX(content=content, pages=pages, overflow=overflow)


__all__ = ["DOCX_FONT", "ResumeDOCX", "build_resume_docx"]

# SPDX-License-Identifier: Apache-2.0
"""Word 渲染的段落级拼装件（v2）：页眉 / 正文块 / lead-trail 列 / Inline 段 -> run。

docx_document_render.py 贴着 500 行预算，这里放与"分区/条目语义"无关的通用件：
页头、正文块、灰字段列、行内段拆 run、角色样式取值。两个模块共享同一套字号/颜色口径
（全部来自 ResumeLayout 与 DocxTheme，不存第二份数值）。
"""

from __future__ import annotations

from collections.abc import Callable, Sequence
from functools import partial
from io import BytesIO

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm

from ..schemas.resume_document import EntryLayout
from ..schemas.resume_document_plan import BlockPlan, CellPlan, DocumentPlan, Inline
from .docx_blocks import (
    DocxTheme,
    decorate_paragraph,
    float_picture,
    meta_line,
    paragraph,
    set_run_font,
)
from .docx_rich import add_rich_runs
from .pdf_exporter import _PX_TO_MM, _PX_TO_PT, _photo_bytes, fit_photo_size
from .resume.resume_contact import WORD_SEPARATOR
from .resume.resume_style import RoleStyle

#: 角色/对象覆盖查询函数（与 PDF / HTML 同一形态）；None = 模板默认。
RoleStyleQuery = Callable[..., RoleStyle | None]

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def plain_text(segments: Sequence[Inline]) -> str:
    """行内段拼成纯文本（条目头的左右两段不再逐段换样式）。"""
    return "".join(str(getattr(seg, "text", "") or "") for seg in segments).strip()


def rich_text(segments: Sequence[Inline]) -> str:
    """行内段拼成带行内标记的文本（strong 段用加粗标记，交给 add_rich_runs 拆 run）。"""
    parts: list[str] = []
    for segment in segments:
        text = str(getattr(segment, "text", "") or "")
        if not text:
            continue
        if getattr(segment, "strong", False) and not text.startswith("**"):
            text = "**" + text + "**"
        parts.append(text)
    return "".join(parts)


def style_of(
    role_style: RoleStyleQuery | None, role: str, object_id: str | None = None
) -> RoleStyle | None:
    """角色/对象覆盖查询；无覆盖返回 None（= 模板值）。"""
    return role_style(role, object_id=object_id) if role_style is not None else None


def scale_of(style: RoleStyle | None) -> float:
    return float(style.font_scale) if style is not None and style.font_scale else 1.0


def color_of(style: RoleStyle | None, fallback: tuple[int, int, int]) -> tuple[int, int, int]:
    """角色覆盖的字色（#rrggbb）；未设或非法回落到模板色。"""
    value = style.color if style is not None else None
    if not isinstance(value, str):
        return fallback
    text = value.strip().lstrip("#")
    if len(text) == 3:
        text = "".join(ch * 2 for ch in text)
    if len(text) != 6:
        return fallback
    try:
        return (int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16))
    except ValueError:
        return fallback


def role_ratio(layout, role: str) -> float:
    """角色 -> 该角色在模板里的字号基准（与 HTML/PDF 的 data-role 层级同源）。"""
    ratios = {
        "name": layout.name_ratio,
        "section_title": layout.section_title_ratio,
        "entry_title": layout.entry_title_ratio,
        "meta": layout.entry_meta_ratio,
        "body": 1.0,
    }
    return float(ratios.get(role, 1.0))


def add_segments(
    node,
    segments: Sequence[Inline],
    *,
    base: float,
    layout,
    theme: DocxTheme,
    role_style: RoleStyleQuery | None,
    default_role: str,
    object_id: str | None = None,
    bold_default: bool = False,
) -> None:
    """把 plan 的 Inline 段写成一段多 run 文本；每段按自己的 role 取字号/颜色。"""
    set_font = partial(set_run_font, font_name=theme.font_name)
    for segment in segments:
        text = str(getattr(segment, "text", "") or "")
        if not text:
            continue
        role = str(getattr(segment, "role", "") or default_role)
        style = style_of(role_style, role, object_id)
        size = base * role_ratio(layout, role) * scale_of(style)
        fallback = theme.body_text if role == "body" else theme.muted_text
        # Q11：run 吃角色/对象覆盖的加粗、斜体与字距——此前只认段内的 ** ** / ❗标记，
        # 面板里调「正文 → 加粗 / 斜体 / 字距」对 Word 完全无效。
        style_bold = style.bold if style is not None and style.bold is not None else None
        bold = bool(getattr(segment, "strong", False)) or (
            style_bold if style_bold is not None else bold_default
        )
        italic = bool(getattr(segment, "italic", False)) or bool(
            style.italic if style is not None and style.italic is not None else False
        )
        letter_spacing_pt = (
            base * float(style.letter_spacing) * _PX_TO_PT
            if style is not None and style.letter_spacing is not None
            else None
        )
        add_rich_runs(
            node,
            text,
            set_font,
            bold=bold,
            italic=italic,
            size_pt=size * _PX_TO_PT,
            color=color_of(style, fallback),
            letter_spacing_pt=letter_spacing_pt,
        )


def draw_header(
    doc,
    plan: DocumentPlan,
    *,
    layout,
    base: float,
    include_photo: bool,
    theme: DocxTheme,
    accent: tuple[int, int, int],
    accent_soft: tuple[int, int, int],
) -> None:
    """页眉：姓名 + 姓名附带字段 + 意向 + 联系方式（planner 已给好每行）+ 浮动照片。"""
    set_font = partial(set_run_font, font_name=theme.font_name)
    name_text = plain_text(plan.header.name)
    meta_text = plain_text(plan.header.name_meta)
    subtitle_text = plain_text(plan.header.subtitle)

    name_node = paragraph(doc)
    header_paragraphs = [name_node]
    if name_text:
        add_rich_runs(
            name_node,
            name_text,
            set_font,
            bold=True,
            size_pt=base * layout.name_ratio * _PX_TO_PT,
            color=accent,
        )
    if meta_text:
        add_rich_runs(
            name_node,
            "　" + meta_text,
            set_font,
            size_pt=base * layout.name_extra_ratio * _PX_TO_PT,
            color=theme.muted_text,
        )

    photo = _photo_bytes(plan.header.photo) if (plan.header.photo and include_photo) else None
    if photo:
        try:
            photo_w = layout.gap_mm(layout.photo_width_ratio)
            photo_h = layout.gap_mm(layout.photo_height_ratio)
            draw_w, draw_h = fit_photo_size(photo, photo_w, photo_h)
            picture_run = name_node.add_run()
            picture_run.add_picture(BytesIO(photo), width=Mm(draw_w))
            float_picture(
                picture_run,
                align="left" if layout.header_photo_side == "left" else "right",
                vertical_offset_mm=max(0.0, (photo_h - draw_h) / 2),
            )
        except Exception:  # noqa: BLE001 - 照片坏了不能阻断导出
            pass

    if subtitle_text:
        subtitle_node = paragraph(doc)
        header_paragraphs.append(subtitle_node)
        add_rich_runs(
            subtitle_node,
            subtitle_text,
            set_font,
            size_pt=base * layout.intent_ratio * _PX_TO_PT,
            color=theme.muted_text,
        )

    contact_lines = [
        WORD_SEPARATOR.join(
            str(seg.text or "").strip() for seg in line if str(seg.text or "").strip()
        )
        for line in plan.header.contact_lines
    ]
    contact_lines = [line for line in contact_lines if line]
    if contact_lines:
        contact_node = paragraph(doc)
        header_paragraphs.append(contact_node)
        for index, line in enumerate(contact_lines):
            if index:
                contact_node.add_run().add_break()
            add_rich_runs(
                contact_node,
                line,
                set_font,
                size_pt=base * layout.contact_ratio * _PX_TO_PT,
                color=theme.muted_text,
            )

    # 页头底色块（modern 的浅色块）：三段各加同色底纹，段距为 0 时连成一整块。
    if getattr(layout, "header_bg", "none") == "soft_accent":
        for node in header_paragraphs:
            decorate_paragraph(
                node,
                fill=accent_soft,
                edges=("top", "left", "bottom", "right"),
                border_color=accent_soft,
                border_sz=4,
                space_pt=6,
            )


def draw_blocks(
    doc,
    blocks: Sequence[BlockPlan],
    *,
    base: float,
    layout,
    theme: DocxTheme,
    role_style: RoleStyleQuery | None,
    object_id: str | None,
) -> None:
    """正文块：连续要点合并成一次 bullets；段落逐段写。"""
    from .docx_blocks import bullets

    body_style = style_of(role_style, "body", object_id)
    pending: list[str] = []

    def flush() -> None:
        if not pending:
            return
        bullets(
            doc,
            list(pending),
            base=base,
            indent_mm=layout.gap_mm(layout.li_indent),
            theme=theme,
            style=body_style,
        )
        pending.clear()

    for block in blocks:
        text = rich_text(block.segments).strip()
        if not text:
            continue
        if block.kind == "bullet":
            pending.append(text)
            continue
        flush()
        # 段落块自身不跨页（与 v1 摘要的 w:keepLines 同口径；Word 没有 CSS 的 break-inside）。
        # Q11：段落缩进与正文对齐此前被 Word 忽略（缩进固定 0、对齐永远左对齐）。
        indent_mm = (
            base * _PX_TO_MM * float(body_style.indent)
            if body_style is not None and body_style.indent is not None
            else 0.0
        )
        node = paragraph(doc, indent_mm=indent_mm, keep_together=True)
        if body_style is not None and body_style.align is not None:
            node.paragraph_format.alignment = align_value(body_style.align)
        if body_style is not None and body_style.line_height:
            node.paragraph_format.line_spacing = float(body_style.line_height)
        add_segments(
            node,
            block.segments,
            base=base,
            layout=layout,
            theme=theme,
            role_style=role_style,
            default_role="body",
            object_id=object_id,
        )
    flush()


def draw_cell_rows(
    doc,
    rows: Sequence[Sequence[CellPlan]],
    layout: EntryLayout,
    *,
    base: float,
    page_layout,
    theme: DocxTheme,
    role_style: RoleStyleQuery | None,
    object_id: str | None,
) -> None:
    """lead / trail：equal = 无框表格（solo 行合并）；flow = 一行段落。"""
    if not rows:
        return
    meta_style = style_of(role_style, "meta", object_id)
    align = layout.align or "equal"
    columns = int(layout.columns or 1)
    if align == "flow" or columns <= 1:
        for row in rows:
            if not row:
                continue
            meta_line(
                doc,
                "".join(rich_text(cell.segments) for cell in row),
                base=base,
                layout=page_layout,
                theme=theme,
                style=meta_style,
            )
        return
    table = doc.add_table(rows=len(rows), cols=columns)
    for row_index, row in enumerate(rows):
        cells = [table.cell(row_index, index) for index in range(columns)]
        if len(row) == 1 and columns > 1:
            cells = [cells[0].merge(cells[-1])]
        for index, cell in enumerate(cells):
            node = cell.paragraphs[0]
            node.paragraph_format.space_after = Mm(0)
            segments = row[min(index, len(row) - 1)].segments if row else []
            add_segments(
                node,
                segments,
                base=base,
                layout=page_layout,
                theme=theme,
                role_style=role_style,
                default_role="meta",
                object_id=object_id,
            )
    # 表格与下一条内容之间留一点段后距（Word 表格没有 margin-bottom）。
    paragraph(doc, space_after_mm=0)


def align_value(value: str | None):
    """角色 align -> Word 段落对齐常量；未知/空返回 None（保持默认）。"""
    return _ALIGN.get(str(value or ""))


__all__ = [
    "RoleStyleQuery",
    "add_segments",
    "align_value",
    "color_of",
    "draw_blocks",
    "draw_cell_rows",
    "draw_header",
    "plain_text",
    "rich_text",
    "role_ratio",
    "scale_of",
    "style_of",
]
